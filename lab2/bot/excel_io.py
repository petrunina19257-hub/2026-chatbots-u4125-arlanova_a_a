"""Экспорт задач в Excel и импорт задач из Excel/CSV.

/export — бот присылает файл tasks.xlsx со всеми задачами чата.
Импорт — пользователь присылает боту .xlsx или .csv, бот добавляет задачи.
Для импорта в файле нужна колонка с текстом задачи («задача», «task» или «text»),
колонки «приоритет»/«priority» и «дедлайн»/«due» — по желанию.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime, tzinfo
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import from_db
from task_parser import PRIORITY_TITLE

MAX_IMPORT_ROWS = 200

EXPORT_HEADERS = ["№", "Создана", "Чат", "Автор", "Задача", "Приоритет", "Исполнитель", "Дедлайн", "Статус", "Выполнена"]
STATUS_TITLE = {"open": "в работе", "done": "выполнена"}
PRIORITY_FILL = {
    "high": PatternFill("solid", fgColor="F8D7DA"),
    "medium": PatternFill("solid", fgColor="FFF3CD"),
    "low": PatternFill("solid", fgColor="D4EDDA"),
}
COLUMN_WIDTHS = [6, 17, 18, 18, 50, 11, 16, 17, 11, 17]

# Как могут называться колонки во входном файле
TEXT_COLUMNS = {"задача", "task", "text", "текст", "название"}
PRIORITY_COLUMNS = {"приоритет", "priority"}
DUE_COLUMNS = {"дедлайн", "due", "срок", "deadline", "дата"}


class ImportFileError(Exception):
    """Файл не получилось прочитать — текст ошибки показываем пользователю."""


@dataclass(frozen=True)
class ImportedRow:
    text: str
    priority: str   # как написано в файле, может быть пустым
    due: str        # дедлайн строкой, например "20.10.2026 15:00"


def _local(value: Optional[str], tz: tzinfo) -> str:
    dt = from_db(value)
    return dt.astimezone(tz).strftime("%d.%m.%Y %H:%M") if dt else ""


def task_values(row, chat_title: str, tz: tzinfo) -> List[str]:
    """Одна задача в виде списка строк — для Google Таблицы и Excel."""
    due = from_db(row["due_at"])
    if due:
        due_text = due.astimezone(tz).strftime("%d.%m.%Y %H:%M" if row["has_time"] else "%d.%m.%Y")
    else:
        due_text = ""
    return [
        str(row["id"]),
        _local(row["created_at"], tz),
        chat_title,
        row["author"] or "",
        row["text"],
        PRIORITY_TITLE.get(row["priority"], row["priority"]),
        row["assignee"] or "",
        due_text,
        STATUS_TITLE.get(row["status"], row["status"]),
        _local(row["done_at"], tz),
    ]


def export_xlsx(rows, chat_title: str, tz: tzinfo) -> io.BytesIO:
    """Создаёт Excel-файл в памяти (на диск ничего не пишем)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Задачи"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4A5568")
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append(task_values(row, chat_title, tz))
        excel_row = sheet.max_row
        sheet.cell(excel_row, 6).fill = PRIORITY_FILL.get(row["priority"], PatternFill())
        if row["status"] == "done":
            for cell in sheet[excel_row]:
                cell.font = Font(color="888888", strike=cell.column == 5)

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions  # фильтры в заголовках, как в «умной таблице»

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        # Excel хранит даты как datetime; 00:00 считаем «без времени»
        return value.strftime("%d.%m.%Y" if (value.hour, value.minute) == (0, 0) else "%d.%m.%Y %H:%M")
    return str(value).strip()


def _read_rows(data: bytes, filename: str) -> List[List[str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as error:
            raise ImportFileError("Не получилось открыть Excel-файл. Проверь, что это .xlsx.") from error
        sheet = workbook.worksheets[0]
        return [[_cell_to_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if name.endswith(".csv"):
        # Excel в русской локали сохраняет CSV в cp1251 и с «;» — поддерживаем оба варианта
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("Не получилось прочитать кодировку CSV.")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise ImportFileError("Поддерживаются только файлы .xlsx и .csv")


def _column(headers: List[str], names: set) -> Optional[int]:
    for index, header in enumerate(headers):
        if header.strip().lower() in names:
            return index
    return None


def read_import_file(data: bytes, filename: str) -> List[ImportedRow]:
    """Читает файл и возвращает задачи для импорта."""
    rows = [row for row in _read_rows(data, filename) if any(row)]
    if not rows:
        raise ImportFileError("Файл пустой.")

    headers = [h.lower() for h in rows[0]]
    text_col = _column(headers, TEXT_COLUMNS)
    if text_col is None:
        raise ImportFileError(
            "Не нашёл колонку с задачами. Назови первую строку-заголовок: задача, приоритет, дедлайн."
        )
    priority_col = _column(headers, PRIORITY_COLUMNS)
    due_col = _column(headers, DUE_COLUMNS)

    def get(row: List[str], col: Optional[int]) -> str:
        return row[col] if col is not None and col < len(row) else ""

    result = []
    for row in rows[1:]:
        text = get(row, text_col)
        if text:
            result.append(ImportedRow(text=text[:500], priority=get(row, priority_col), due=get(row, due_col)))
    if not result:
        raise ImportFileError("В файле нет ни одной задачи.")
    if len(result) > MAX_IMPORT_ROWS:
        raise ImportFileError(f"Слишком много строк: {len(result)}. Максимум — {MAX_IMPORT_ROWS}.")
    return result
