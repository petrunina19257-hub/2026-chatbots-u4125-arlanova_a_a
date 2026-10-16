"""Автотесты: разбор задач, база, Excel/CSV, Google Таблица, отзывы. Запуск: python -m unittest -v"""
import os
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import sqlite3

from openpyxl import Workbook, load_workbook

from database import SCHEMA, TaskStorage
from excel_io import ImportFileError, export_xlsx, read_import_file, task_values
from sheets import HEADERS, SheetSync
from task_parser import compute_remind_at, format_due, parse_task

TZ = ZoneInfo("Europe/Moscow")
# Пятница, 16 октября 2026, 12:00
NOW = datetime(2026, 10, 16, 12, 0, tzinfo=TZ)


class ParserTests(unittest.TestCase):
    def test_priority(self):
        self.assertEqual(parse_task("Срочно отправить отчёт", NOW).priority, "high")
        self.assertEqual(parse_task("!! позвонить", NOW).priority, "high")
        self.assertEqual(parse_task("Не срочно почитать статью", NOW).priority, "low")
        self.assertEqual(parse_task("Купить молоко", NOW).priority, "medium")

    def test_text_cleanup(self):
        self.assertEqual(parse_task("!! позвонить #срочно", NOW).text, "Позвонить")
        self.assertEqual(parse_task("Срочно отправить отчёт завтра в 15:00", NOW).text, "Отправить отчёт")
        self.assertEqual(parse_task("Купить подарок до пятницы", NOW).text, "Купить подарок")
        self.assertEqual(parse_task("Позвонить в банк через 2 часа", NOW).text, "Позвонить в банк")
        self.assertEqual(parse_task("Сдать лабу до 20.10, дедлайн жёсткий", NOW).text, "Сдать лабу, дедлайн жёсткий")
        self.assertEqual(parse_task("Завтра", NOW).text, "Завтра")  # пустой текст не оставляем

    def test_tomorrow_with_time(self):
        due = parse_task("Отчёт завтра 15:00", NOW).due
        self.assertEqual(due.at, datetime(2026, 10, 17, 15, 0, tzinfo=TZ))
        self.assertTrue(due.has_time)

    def test_day_after_tomorrow(self):
        due = parse_task("Сдать лабу послезавтра", NOW).due
        self.assertEqual(due.at.date(), datetime(2026, 10, 18).date())
        self.assertFalse(due.has_time)

    def test_weekday(self):
        # сегодня пятница -> «до пятницы» значит следующая пятница
        self.assertEqual(parse_task("до пятницы", NOW).due.at.day, 23)
        self.assertEqual(parse_task("в среду", NOW).due.at.day, 21)
        self.assertEqual(parse_task("к понедельнику", NOW).due.at.day, 19)

    def test_numeric_and_word_dates(self):
        self.assertEqual(parse_task("до 20.10", NOW).due.at.date(), datetime(2026, 10, 20).date())
        self.assertEqual(parse_task("до 01.02", NOW).due.at.year, 2027)  # дата прошла -> след. год
        self.assertEqual(parse_task("5 ноября", NOW).due.at.month, 11)
        self.assertIsNone(parse_task("31.02", NOW).due)

    def test_relative(self):
        self.assertEqual(parse_task("через 2 часа", NOW).due.at, NOW + timedelta(hours=2))
        self.assertEqual(parse_task("через 30 минут", NOW).due.at, NOW + timedelta(minutes=30))
        self.assertEqual(parse_task("через неделю", NOW).due.at, NOW + timedelta(weeks=1))

    def test_time_only(self):
        self.assertEqual(parse_task("в 18:30", NOW).due.at, NOW.replace(hour=18, minute=30))
        self.assertEqual(parse_task("в 09:00", NOW).due.at.day, 17)  # уже прошло -> завтра

    def test_assignee(self):
        self.assertEqual(parse_task("Смета для @alina_k до пятницы", NOW).assignee, "@alina_k")

    def test_no_due(self):
        self.assertIsNone(parse_task("Купить молоко", NOW).due)

    def test_remind(self):
        due = parse_task("завтра 15:00", NOW).due
        self.assertEqual(compute_remind_at(due, NOW), due.at - timedelta(hours=1))
        soon = parse_task("через 30 минут", NOW).due
        self.assertEqual(compute_remind_at(soon, NOW), soon.at)
        date_only = parse_task("завтра", NOW).due
        self.assertEqual(compute_remind_at(date_only, NOW).hour, 10)

    def test_format_due(self):
        self.assertEqual(format_due(NOW.replace(hour=18), True, NOW), "сегодня 18:00")
        self.assertEqual(format_due(NOW + timedelta(days=1), False, NOW), "завтра")
        self.assertEqual(format_due(NOW + timedelta(days=4), False, NOW), "вт 20.10")


class StorageTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = TaskStorage(os.path.join(self.tmp.name, "test.db"))

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, text, priority="medium", due=None, remind=None, chat_id=1):
        return self.db.add_task(chat_id, 10, "Тест", text, priority, "", due, due is not None, remind)

    def test_sorting_and_done(self):
        low = self.add("низкий", "low")
        high = self.add("высокий", "high")
        self.add("средний")
        self.assertEqual([r["id"] for r in self.db.list_open(1)][0], high)
        self.assertTrue(self.db.mark_done(1, high))
        self.assertFalse(self.db.mark_done(1, high))  # повторно закрыть нельзя
        self.assertEqual(len(self.db.list_open(1)), 2)
        self.assertFalse(self.db.mark_done(2, low))  # чужой чат

    def test_reminders(self):
        now = datetime.now(timezone.utc)
        task_id = self.add("напомни", due=now, remind=now - timedelta(minutes=1))
        self.add("позже", due=now, remind=now + timedelta(hours=1))
        self.assertEqual([r["id"] for r in self.db.due_reminders(now)], [task_id])
        self.db.mark_reminded(task_id)
        self.assertEqual(self.db.due_reminders(now), [])

    def test_digest_settings(self):
        self.add("задача")
        self.assertEqual(self.db.digest_chats(), [1])
        self.db.set_digest(1, False)
        self.assertEqual(self.db.digest_chats(), [])


class DataTests(unittest.TestCase):
    """Новое в лабораторной 2."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "test.db")
        self.db = TaskStorage(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, text, priority="medium", chat_id=1):
        return self.db.add_task(chat_id, 10, "Тест", text, priority, "@alina_k", None, False, None, "Личка")

    def test_search_is_case_insensitive_for_russian(self):
        self.add("Подготовить ОТЧЁТ")
        self.add("Купить молоко")
        self.assertEqual(len(self.db.search(1, "отчёт")), 1)
        self.assertEqual(len(self.db.search(1, "ALINA")), 2)  # по исполнителю
        self.assertEqual(self.db.search(2, "отчёт"), [])

    def test_stats(self):
        first = self.add("а", "high")
        self.add("б", "low")
        self.db.mark_done(1, first)
        s = self.db.stats(1, datetime.now(timezone.utc) - timedelta(days=7))
        self.assertEqual((s["total"], s["open"], s["done"], s["low"], s["done_week"]), (2, 1, 1, 1, 1))

    def test_migration_from_lab1_database(self):
        path = os.path.join(self.tmp.name, "old.db")
        # схема базы из лабораторной 1 — без колонки chat_title
        lab1_schema = "\n".join(line for line in SCHEMA.splitlines() if "chat_title" not in line)
        with sqlite3.connect(path) as conn:
            conn.executescript(lab1_schema)
        TaskStorage(path)  # не должно упасть
        with sqlite3.connect(path) as conn:
            columns = [row[1] for row in conn.execute("PRAGMA table_info(tasks)")]
        self.assertIn("chat_title", columns)

    def test_export_xlsx(self):
        self.add("Задача <1>", "high")
        buffer = export_xlsx(self.db.all_tasks(1), "Личка", TZ)
        sheet = load_workbook(buffer).active
        self.assertEqual(sheet["E2"].value, "Задача <1>")
        self.assertEqual(sheet["F2"].value, "высокий")
        self.assertEqual(sheet.max_row, 2)

    def test_import_xlsx(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Задача", "Приоритет", "Дедлайн"])
        sheet.append(["Сдать лабу", "высокий", datetime(2026, 10, 20, 15, 0)])
        sheet.append(["Купить хлеб", None, None])
        sheet.append([None, None, None])  # пустая строка пропускается
        path = os.path.join(self.tmp.name, "in.xlsx")
        workbook.save(path)
        with open(path, "rb") as f:
            rows = read_import_file(f.read(), "in.xlsx")
        self.assertEqual(len(rows), 2)
        self.assertEqual((rows[0].priority, rows[0].due), ("высокий", "20.10.2026 15:00"))

    def test_import_csv_cp1251_semicolon(self):
        data = "задача;дедлайн\nПозвонить маме;завтра\n".encode("cp1251")
        rows = read_import_file(data, "tasks.csv")
        self.assertEqual((rows[0].text, rows[0].due), ("Позвонить маме", "завтра"))

    def test_import_errors(self):
        with self.assertRaises(ImportFileError):
            read_import_file(b"name;email\n1;2", "x.csv")  # нет колонки «задача»
        with self.assertRaises(ImportFileError):
            read_import_file(b"hello", "x.pdf")
        with self.assertRaises(ImportFileError):
            read_import_file(b"not a zip", "x.xlsx")

    def test_sheet_sync_upsert_and_delete(self):
        sheet = FakeWorksheet()
        sync = SheetSync(sheet)
        sync.ensure_headers()
        task_id = self.add("Синхронизация")
        values = task_values(self.db.get_task(1, task_id), "Личка", TZ)
        sync.upsert(values)
        self.db.mark_done(1, task_id)
        sync.upsert(task_values(self.db.get_task(1, task_id), "Личка", TZ))
        self.assertEqual(len(sheet.rows), 2)  # строка обновилась, а не задвоилась
        self.assertEqual(sheet.rows[1][8], "выполнена")
        sync.delete(task_id)
        self.assertEqual(sheet.rows, [HEADERS])


class FeedbackTests(unittest.TestCase):
    """Новое в лабораторной 3."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = TaskStorage(os.path.join(self.tmp.name, "test.db"))

    def tearDown(self):
        self.tmp.cleanup()

    def test_rating_then_text_is_one_record(self):
        feedback_id = self.db.add_feedback(42, "@alina_k", 4, None)
        row = self.db.add_feedback_text(feedback_id, "Удобно")
        self.assertEqual((row["rating"], row["text"]), (4, "Удобно"))
        self.assertEqual(len(self.db.list_feedback()), 1)

    def test_usage_stats(self):
        self.db.log_event(1, 1, "/list")
        self.db.log_event(1, 1, "/list")
        self.db.log_event(2, 2, "text")
        self.db.add_feedback(1, "a", 5, None)
        self.db.add_feedback(2, "b", 3, "норм")
        self.db.add_feedback(3, "c", None, "без оценки")
        s = self.db.usage_stats(datetime.now(timezone.utc) - timedelta(days=1))
        self.assertEqual((s["users"], s["active_day"], s["events"]), (2, 2, 3))
        self.assertEqual(s["actions"][0], ("/list", 2))
        self.assertEqual((s["feedback"], s["rated"], s["avg_rating"]), (3, 2, 4.0))


class FakeWorksheet:
    """Подмена листа Google Таблицы для тестов без интернета."""

    class _Spreadsheet:
        url = "https://docs.google.com/spreadsheets/d/test"

    spreadsheet = _Spreadsheet()

    def __init__(self):
        self.rows = []

    def row_values(self, n):
        return self.rows[n - 1] if len(self.rows) >= n else []

    def col_values(self, n):
        return [row[n - 1] for row in self.rows]

    def update(self, range_name, values, raw=True):
        start = int(range_name[1:]) - 1
        for offset, row in enumerate(values):
            if start + offset < len(self.rows):
                self.rows[start + offset] = list(row)
            else:
                self.rows.append(list(row))

    def freeze(self, rows):
        pass

    def append_row(self, values, value_input_option="RAW"):
        self.rows.append(list(values))

    def delete_rows(self, n):
        del self.rows[n - 1]

    def clear(self):
        self.rows = []


class AppTests(unittest.TestCase):
    def test_build_application(self):
        os.environ["BOT_TOKEN"] = "123456:TEST"
        with tempfile.TemporaryDirectory() as tmp:
            os.environ["DB_PATH"] = os.path.join(tmp, "app.db")
            os.environ.pop("GOOGLE_SHEET_ID", None)
            import bot
            app = bot.build_application()
            self.assertGreater(len(app.handlers[0]), 20)
            self.assertEqual(len(app.handlers[-1]), 1)  # счётчик статистики видит всё первым
            self.assertIsNone(app.bot_data["sheet"])  # без настроек бот работает без таблицы
            self.assertEqual(len(app.job_queue.jobs()), 2)  # напоминания + дайджест


if __name__ == "__main__":
    unittest.main()
